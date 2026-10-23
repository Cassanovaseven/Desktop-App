import struct
import zipfile
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches


def minimal_pdf(text: str) -> bytes:
    """PDF satu halaman dengan teks Helvetica, ditulis manual supaya tidak butuh pustaka pembuat PDF."""
    stream = f"BT /F1 24 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def minimal_psd(width: int, height: int, rgb: tuple[int, int, int]) -> bytes:
    """PSD RGB 8-bit tanpa lapisan, data gambar gabungan tanpa kompresi (Pillow tidak bisa menulis PSD)."""
    header = b"8BPS" + struct.pack(">H6xHIIHH", 1, 3, height, width, 8, 3)
    sections = struct.pack(">III", 0, 0, 0)  # color mode data, image resources, layer & mask
    planes = b"".join(bytes([c]) * (width * height) for c in rgb)
    return header + sections + struct.pack(">H", 0) + planes


@pytest.fixture
def vault(tmp_path: Path) -> Path:
    root = tmp_path / "Vault Alfonsó"
    for d in ("files/ab", "thumbs"):
        (root / d).mkdir(parents=True)
    return root


@pytest.fixture
def make(vault: Path):
    """Membuat file contoh kecil di vault/files/ab/ dan mengembalikan path relatifnya."""

    def _make(kind: str) -> str:
        d = vault / "files" / "ab"
        if kind == "pdf":
            (d / "c.pdf").write_bytes(minimal_pdf("pertanggungjawaban pidana"))
            return "files/ab/c.pdf"
        if kind == "pdf-rusak":
            (d / "r.pdf").write_bytes(b"%PDF-1.4 bukan pdf")
            return "files/ab/r.pdf"
        if kind == "docx":
            doc = Document()
            doc.add_heading("Bab Satu", level=1)
            doc.add_paragraph("Isi tentang riset hukum.")
            doc.save(d / "c.docx")
            return "files/ab/c.docx"
        if kind == "pptx":
            prs = Presentation()
            s = prs.slides.add_slide(prs.slide_layouts[5])
            s.shapes.title.text = "Judul Slide"
            box = s.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
            box.text_frame.text = "poin presentasi"
            prs.save(d / "c.pptx")
            return "files/ab/c.pptx"
        if kind == "xlsx":
            wb = Workbook()
            wb.active.append(["nama", "nilai"])
            wb.active.append(["anggaran", 1200])
            wb.save(d / "c.xlsx")
            return "files/ab/c.xlsx"
        if kind == "csv":
            (d / "c.csv").write_text("kota,jumlah\nBandung,3\n", encoding="utf-8")
            return "files/ab/c.csv"
        if kind == "md":
            (d / "c.md").write_text("# Catatan\nkata kunci rapat", encoding="utf-8")
            return "files/ab/c.md"
        if kind == "html":
            (d / "c.html").write_text("<html><head><style>p{}</style><script>x=1</script></head><body><p>halo <b>dunia</b></p></body></html>", encoding="utf-8")
            return "files/ab/c.html"
        if kind == "epub":
            with zipfile.ZipFile(d / "c.epub", "w") as z:
                z.writestr("mimetype", "application/epub+zip")
                z.writestr("OEBPS/bab1.xhtml", "<html><body><h1>Bab</h1><p>cerita panjang</p></body></html>")
            return "files/ab/c.epub"
        if kind == "py":
            (d / "c.py").write_text("def halo():\n    return 'dunia'\n", encoding="utf-8")
            return "files/ab/c.py"
        if kind == "png":
            Image.new("RGB", (800, 600), (200, 30, 30)).save(d / "c.png")
            return "files/ab/c.png"
        if kind == "png-teks":
            img = Image.new("RGB", (640, 120), "white")
            ImageDraw.Draw(img).text((20, 30), "Laporan pidana korporasi", fill="black", font=ImageFont.load_default(size=40))
            img.save(d / "t.png")
            return "files/ab/t.png"
        if kind == "png-dominan":
            # Sebagian besar merah, sebagian kecil biru: warna dominan = merah, bukan rata-ratanya.
            img = Image.new("RGB", (100, 100), (220, 20, 20))
            img.paste((20, 20, 220), (0, 0, 100, 30))
            img.save(d / "dom.png")
            return "files/ab/dom.png"
        if kind == "gif":
            frames = [Image.new("RGB", (40, 30), c) for c in ((0, 200, 0), (0, 0, 200))]
            frames[0].save(d / "a.gif", save_all=True, append_images=frames[1:], duration=100, loop=0)
            return "files/ab/a.gif"
        if kind == "psd":
            (d / "c.psd").write_bytes(minimal_psd(32, 24, (10, 120, 200)))
            return "files/ab/c.psd"
        if kind == "psd-rusak":
            (d / "r.psd").write_bytes(b"8BPS\x00\x01rusak")
            return "files/ab/r.psd"
        if kind == "zip":
            with zipfile.ZipFile(d / "c.zip", "w") as z:
                z.writestr("catatan.txt", "isi catatan")
                z.writestr("sub/gambar.md", "# md")
                z.writestr("../jahat.txt", "keluar")
                z.writestr("/absolut.txt", "keluar")
            return "files/ab/c.zip"
        if kind == "zip-bom":
            with zipfile.ZipFile(d / "bom.zip", "w") as z:
                for i in range(2001):
                    z.writestr(f"f{i}.txt", "x")
            return "files/ab/bom.zip"
        if kind == "bin":
            (d / "c.xyz").write_bytes(b"\x00\x01\x02")
            return "files/ab/c.xyz"
        raise ValueError(kind)

    return _make
