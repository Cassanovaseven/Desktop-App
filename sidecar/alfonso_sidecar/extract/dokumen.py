import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from .teks import strip_html


class PdfError(Exception):
    pass


class OfficeError(Exception):
    pass


def pdf(path: Path):
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise PdfError("terenkripsi")
        pages = [page.extract_text() or "" for page in reader.pages]
    except PdfError:
        raise
    except Exception as e:
        raise PdfError(str(e)) from e
    return "\n".join(pages), {"halaman": len(pages)}


def docx(path: Path):
    try:
        doc = Document(str(path))
    except Exception as e:
        raise OfficeError(str(e)) from e
    paras = [p for p in doc.paragraphs if p.text.strip()]
    headings = [p.text for p in paras if p.style is not None and p.style.name.lower().startswith("heading")]
    return "\n".join(p.text for p in paras), {"heading": headings}


def pptx(path: Path):
    try:
        prs = Presentation(str(path))
    except Exception as e:
        raise OfficeError(str(e)) from e
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
    return "\n".join(parts), {"slide": len(prs.slides)}


def xlsx(path: Path):
    try:
        wb = load_workbook(str(path), read_only=True, data_only=True)
    except Exception as e:
        raise OfficeError(str(e)) from e
    lines = []
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(v) for v in row if v is not None]
                if cells:
                    lines.append("\t".join(cells))
        sheets = len(wb.worksheets)
    finally:
        wb.close()
    return "\n".join(lines), {"sheet": sheets}


def epub(path: Path):
    try:
        with zipfile.ZipFile(path) as z:
            names = sorted(n for n in z.namelist() if n.lower().endswith((".xhtml", ".html", ".htm")))
            parts = [strip_html(z.read(n).decode("utf-8", errors="replace")) for n in names]
    except zipfile.BadZipFile as e:
        raise OfficeError(str(e)) from e
    return "\n".join(p for p in parts if p), {"bab": len(names)}
